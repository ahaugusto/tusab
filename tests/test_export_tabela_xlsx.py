# Copyright (c) 2026 CriAugu — CNPJ 65.131.075/0001-57
"""
POST /export/tabela-xlsx — planilha a partir das tabelas de uma resposta do
chat (botão "Baixar planilha" no ChatDrawer) + reforço de prompt para pedido
explícito de tabela/planilha (agent/chat.py::_instrucao_pedido_tabela).

Foco: números saem como NÚMERO (pt-BR "1.234,56" e en "1,234.56"),
identificadores ("001", CNPJ) continuam texto, datas só quando inequívocas,
cabeçalho em negrito e congelado, anti-injeção de fórmula, limites de entrada.
"""
import io
import importlib
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

exp = importlib.import_module("tusab_engine.api.router_exports")
chat_mod = importlib.import_module("tusab_engine.agent.chat")


def _tipar(valores, idioma='pt'):
    return exp._tipar_coluna(valores, idioma)


# ── Tipagem de colunas ─────────────────────────────────────────────────────────

def test_numeros_pt_br_com_milhar_e_decimal():
    r = _tipar(['1.234,56', '10,5', '3'])
    assert [v for v, _ in r] == [1234.56, 10.5, 3]
    assert r[0][1] == '#,##0.00'


def test_numeros_en_detectados_pela_forma():
    r = _tipar(['1,234.56', '10.5', '3'])
    assert [v for v, _ in r] == [1234.56, 10.5, 3]


def test_milhar_ambiguo_segue_idioma():
    # "1.234" sozinho: pt → milhar (1234); en → decimal (1.234)
    assert _tipar(['1.234', '2.500'], 'pt')[0][0] == 1234
    assert _tipar(['1.234', '2.500'], 'en')[0][0] == pytest.approx(1.234)
    # evidência de outra linha decide a coluna inteira
    assert [v for v, _ in _tipar(['1.234', '12,5'], 'en')] == [1234, 12.5]


def test_moeda_percentual_e_negativos():
    r = _tipar(['R$ 1.234,56', 'R$ -10,00', '(5,00)'])
    assert [v for v, _ in r] == [1234.56, -10.0, -5.0]
    assert r[0][1].startswith('"R$"')
    p = _tipar(['12,5%', '7%'])
    assert p[0][0] == pytest.approx(0.125) and p[1][0] == pytest.approx(0.07)
    assert p[0][1].endswith('%')


def test_ano_inteiro_sem_formato_de_milhar():
    r = _tipar(['2023', '2024'])
    assert r == [(2023, None), (2024, None)]


def test_identificadores_continuam_texto():
    assert _tipar(['001', '002'])[0] == ('001', None)
    assert _tipar(['65131075000157', '12345678901234567'])[1][0] == '12345678901234567'
    assert _tipar(['12 kg', '3'])[1] == ('3', None)   # coluna mista → texto


def test_placeholders_nao_impedem_coluna_numerica():
    r = _tipar(['10', '-', 'n/d', '', '20'])
    assert r[0] == (10, None) and r[4] == (20, None)
    assert r[1] == ('-', None) and r[3] == (None, None)


def test_datas_inequivocas_e_ambiguas():
    r = _tipar(['15/03/2024', '01/04/2024'], 'pt')
    assert r[0][0] == date(2024, 3, 15) and r[1][0] == date(2024, 4, 1)
    assert r[0][1] == 'dd/mm/yyyy'
    assert _tipar(['2024-03-15'])[0][0] == date(2024, 3, 15)
    # em inglês "03/04/2024" é ambíguo de verdade → fica texto
    assert _tipar(['03/04/2024'], 'en')[0] == ('03/04/2024', None)
    # dia > 12 na segunda posição → mm/dd
    assert _tipar(['03/25/2024'], 'en')[0][0] == date(2024, 3, 25)
    # data inválida → texto
    assert _tipar(['31/02/2024'])[0] == ('31/02/2024', None)


# ── Endpoint ───────────────────────────────────────────────────────────────────

def _payload(**kw):
    base = {
        "titulo": "Compare os preços de 2024",
        "idioma": "pt",
        "tabelas": [{
            "nome": "",
            "cabecalho": ["Produto", "Preço (R$)", "Data", "Código"],
            "linhas": [
                ["Café", "1.234,56", "15/03/2024", "001"],
                ["=HYPERLINK(\"http://x\")", "10,5", "02/04/2024", "002"],
            ],
        }],
    }
    base.update(kw)
    return base


def test_endpoint_gera_planilha_organizada(client):
    res = client.post("/export/tabela-xlsx", json=_payload())
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    hoje = datetime.now().strftime('%Y-%m-%d')
    assert f'tusab_compare-os-precos-de-2024_{hoje}.xlsx' in res.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active
    assert ws.title == "Tabela 1"
    assert ws.freeze_panes == "A2"
    assert all(ws.cell(row=1, column=c).font.bold for c in range(1, 5))
    assert ws["B2"].value == 1234.56 and ws["B2"].data_type == 'n'
    assert ws["B3"].value == 10.5
    assert ws["C2"].is_date and ws["C2"].value.date() == date(2024, 3, 15)
    assert ws["D2"].value == "001" and ws["D2"].data_type == 's'
    # anti-injeção de fórmula: continua string, não vira fórmula
    assert ws["A3"].data_type == 's'
    assert ws["A3"].value.startswith("=HYPERLINK")
    assert ws.column_dimensions["A"].width >= 8
    assert ws.auto_filter.ref == "A1:D3"


def test_endpoint_uma_aba_por_tabela_com_nomes_unicos(client):
    t = {"nome": "Resumo", "cabecalho": ["A"], "linhas": [["1"]]}
    res = client.post("/export/tabela-xlsx", json=_payload(tabelas=[t, t, {"cabecalho": ["B"], "linhas": []}]))
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Resumo", "Resumo (2)", "Tabela 3"]


def test_endpoint_completa_linhas_curtas_e_corta_longas(client):
    res = client.post("/export/tabela-xlsx", json=_payload(tabelas=[
        {"cabecalho": ["A", "B"], "linhas": [["1"], ["2", "3", "4"]]}]))
    ws = load_workbook(io.BytesIO(res.content)).active
    assert ws.max_column == 2
    assert ws["B2"].value is None and ws["B3"].value == 3


@pytest.mark.parametrize("tabelas", [
    [],                                                                  # nenhuma tabela
    [{"cabecalho": [], "linhas": []}],                                   # sem colunas
    [{"cabecalho": ["c"] * 31, "linhas": []}],                           # colunas demais
    [{"cabecalho": ["a"], "linhas": [["x"]] * 2001}],                    # linhas demais
    [{"cabecalho": ["a"], "linhas": [["x" * 2001]]}],                    # célula grande demais
    [{"cabecalho": ["a"] * 30, "linhas": [["x"] * 30] * 700}],           # total de células
    [{"cabecalho": ["a"], "linhas": []}] * 11,                           # tabelas demais
])
def test_endpoint_valida_limites(client, tabelas):
    res = client.post("/export/tabela-xlsx", json=_payload(tabelas=tabelas))
    assert res.status_code == 422


# ── Prompt: pedido explícito de tabela/planilha ────────────────────────────────

@pytest.mark.parametrize("pergunta", [
    "Gere uma planilha com os preços",
    "compare X e Y numa tabela",
    "make a table of the results",
    "exporta pra Excel",
    "coloque em formato de tabela",
    "haz una tabla con los datos",
])
def test_instrucao_pedido_tabela_dispara(pergunta):
    assert "PEDIDO DE TABELA" in chat_mod._instrucao_pedido_tabela(pergunta)


def test_instrucao_pedido_tabela_nao_dispara_em_pergunta_explicativa():
    assert chat_mod._instrucao_pedido_tabela("Explique o que é o SUS") == ""
    # perguntas SOBRE uma tabela não são pedido de formato
    assert chat_mod._instrucao_pedido_tabela("Qual o valor da tabela SUS para consulta?") == ""
    assert chat_mod._instrucao_pedido_tabela("Explique a tabela periódica") == ""
    assert chat_mod._instrucao_pedido_tabela("") == ""


def test_prompt_mantem_anti_injecao_e_regras_de_tabela():
    ctx = [{"titulo": "t", "data": "", "link": "", "texto": "conteúdo"}]
    p = chat_mod._montar_prompt("gere uma tabela com os valores", ctx)
    assert chat_mod._ANTI_INJECAO_INSTR in p
    assert "|---|---|" in p
    assert "PEDIDO DE TABELA" in p
    # pedido explícito troca o bloco de formato geral (listas) pelo só-tabela
    assert chat_mod._FMT_INSTR_TABELA in p and chat_mod._FMT_INSTR not in p
    # instrução de tabela vem antes das fontes/pergunta, não dentro delas
    assert p.index("PEDIDO DE TABELA") < p.index("<sources canal=")
    p2 = chat_mod._montar_prompt("explique o tema", ctx)
    assert "PEDIDO DE TABELA" not in p2
    assert chat_mod._FMT_INSTR in p2 and chat_mod._ANTI_INJECAO_INSTR in p2


def test_prompt_contexto_tambem_usa_formato_de_tabela():
    p = chat_mod._montar_prompt_contexto("transforme isso numa tabela", [], {"pergunta": "q", "resposta": "r"})
    assert chat_mod._FMT_INSTR_TABELA in p and "PEDIDO DE TABELA" in p


def test_persona_preservada_com_pedido_de_tabela():
    ctx = [{"titulo": "t", "data": "", "link": "", "texto": "conteúdo"}]
    p = chat_mod._montar_prompt("gere uma planilha", ctx, persona="didatico")
    assert "TOM DE RESPOSTA" in p
